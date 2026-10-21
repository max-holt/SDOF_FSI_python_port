"""Extract the golden fixtures from the source workbook (``curve_10in`` and ``Example``).

Reads ``tests/fixtures/source/surge_force_data.xlsm`` (Excel-cached values, i.e. the
numbers the VBA UDFs produced inside Excel) and writes, under ``tests/fixtures/curve_10in/``:

* ``force_history.csv`` - the full numeric extent of ``Sheet1`` (time, history_A,
  history_B); the test slices the first ``n_samples`` rows of ``history_B``.
* ``params.json`` - the arguments common to every case, taken from the parameter cells
  of ``curve_10in`` exactly as Excel passed them to the UDFs (cached doubles).
* ``cases.csv`` - one row per live UDF call in the case table (column H array formula
  ``GetPeakResponse`` -> solver ``lumped``, ``GetPeakResponseJC`` -> ``jc``): the
  per-case inputs (E: physical length, F: natural frequency, G: dynamic friction) and
  the five expected outputs H..L. ``source_row`` is the workbook row the case came from.

Rows whose column H holds a pasted number instead of a formula (rows 100-101 at the time
of writing) are not cases and are skipped.

The ``Example`` sheet yields ``tests/fixtures/example/cases.json``: three self-contained
fast cases (inputs, keyword arguments and expected outputs in one record each):

* ``example_pressure`` - ``GetPressureResponse`` on the 5-value valve pattern B3:B7 with
  parameters J33:J42 (``pipeArea = PI()*(J39^2/4)``; Excel's ``PI()`` is the IEEE double
  pi, which was verified to reproduce more cached outputs bit-exactly than the 15-digit
  literal), cached output C3:C123 (121 samples).
* ``example_lumped`` - ``GetPeakResponse`` on the force column E3:E27 with J6:J14, output F3:F7.
* ``example_jc`` - ``GetPeakResponseJC`` on E3:E27 with J19:J29, output G3:G7.

``tests/fixtures/manifest.json`` lists both fixture sets (directory, files, source sheet,
case counts and markers); ``tests/test_golden.py`` reads it. The argument mapping of every formula is
checked against the expected shape so a workbook layout change is caught here, not in
the test.

Floats are written with ``repr`` (shortest round-trip), so the fixtures carry the exact
doubles Excel held.

Run manually (this is the only sanctioned way to produce files under tests/fixtures/):

    python tools/extract_fixtures.py

Requires ``openpyxl`` (a ``dev`` extra: ``pip install -e ".[dev]"``).
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
import sys
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = REPO / "tests" / "fixtures" / "source" / "surge_force_data.xlsm"
DEFAULT_OUT = REPO / "tests" / "fixtures"

CASE_SHEET = "curve_10in"
FORCE_SHEET = "Sheet1"
FORCE_HEADER = ("time", "history_A", "history_B")
# curve_10in!C{r} is "=Sheet1!D{r+1}", and Sheet1 column D is history_B.
FORCE_COLUMN = "history_B"
FORCE_SHEET_COL = 4  # D
FIRST_CASE_ROW = 20

# Parameter cells of curve_10in: (label cell, expected label, value cell, name)
PARAM_CELLS = [
    ("E4", "Step QTY", "H4", "step_qty"),
    ("E5", "Sample Frequency", "H5", "sample_freq"),
    ("E6", "Celerity", "H6", "celerity"),
    ("E9", "Density Fluid", "H9", "fluid_densty"),
    ("E11", "Mass/Physical Length", "H11", "mass_per_unit_len"),
    ("E12", "Damping (% Critical)", "H12", "damping_percent"),
    ("E13", "Pipe Area", "H13", "fluid_area"),
    ("E14", "E", "H14", "youngs_modulus"),
    ("E15", "I", "H15", "second_moment"),
    ("N4", "frequency factor", "P4", "restrain_freq_ratio"),
    ("N5", "gap", "P5", "gap"),
    ("N12", "Force Correction", "N13", "force_correction"),
]

FORCE_ARG = r"\$C\$(\d+):\$C\$(\d+)\*1000\*\$N\$13"
LUMPED_RE = re.compile(
    r"^=GetPeakResponse\(" + FORCE_ARG + r",\s*\$H\$4,\s*\$H\$5,\s*E(\d+),\s*\$H\$14\*\$H\$15,"
    r"\s*\$H\$6,\s*\$H\$11,\s*\$H\$12,\s*\$G(\d+),\s*F(\d+),\s*\$P\$4,\s*\$P\$5\)$"
)
JC_RE = re.compile(
    r"^=GetPeakResponseJC\(" + FORCE_ARG + r",\s*\$H\$4,\s*\$H\$5,\s*E(\d+),\s*\$H\$14\*\$H\$15,"
    r"\s*\$H\$6,\s*\$H\$11,\s*\$H\$9,\s*\$H\$13,\s*\$H\$12,\s*G(\d+),\s*F(\d+),\s*\$P\$4,\s*\$P\$5,\s*0\)$"
)
OUTPUT_NAMES = (
    "max_static_equivalent_force",
    "max_displacement",
    "max_velocity",
    "max_acceleration",
    "max_applied_force",
)

EXAMPLE_SHEET = "Example"
# (formula cell, expected spill range, expected formula with all spaces removed, solver)
EXAMPLE_FORMULAS = [
    (
        "C3",
        "C3:C123",
        "=TRANSPOSE(GetPressureResponse(B3:B7,J33,J34,J35,J36,J37,J38,PI()*(J39^2/4),J40,J41,J42))",
        "pressure",
    ),
    (
        "F3",
        "F3:F7",
        "=TRANSPOSE(GetPeakResponse(E3:E27,J6,J7,J8,J9,J10,J11,J12,J13,J14))",
        "lumped",
    ),
    (
        "G3",
        "G3:G7",
        "=TRANSPOSE(GetPeakResponseJC(E3:E27,J19,J20,J21,J22,J23,J24,J25,J26,J27,J28,J29))",
        "jc",
    ),
]
# (label cell, expected label, value cell, keyword) per Example case, in VBA argument order
EXAMPLE_PARAMS = {
    "pressure": [
        ("I33", "Valve Closure Time", "J33", "closure_time"),
        ("I34", "Celerity", "J34", "celerity"),
        ("I35", "Density", "J35", "density"),
        ("I36", "Time Step", "J36", "time_step"),
        ("I37", "SystemPressure Drop", "J37", "pressure_drop_from_pump"),
        ("I38", "Flowrate", "J38", "flow_rate"),
        ("I39", "Pipe ID", "J39", "pipe_id_mm"),  # pipe_area = PI()*(J39^2/4), see below
        ("I40", "Valve Cv", "J40", "valve_cv"),
        ("I41", "Distance To Pump", "J41", "length_to_pump"),
        ("I42", "Pump Ramp-up", "J42", "pump_ramp_pressure"),
    ],
    "lumped": [
        ("I6", "Step QTY", "J6", "step_qty"),
        ("I7", "Sample Frequency", "J7", "sample_freq"),
        ("I8", "Physical Length", "J8", "physical_length"),
        ("I9", "EI", "J9", "flexural_rigidity_ei"),
        ("I10", "Celerity", "J10", "celerity"),
        ("I11", "Mass/Physical Length", "J11", "mass_per_unit_len"),
        ("I12", "Damping (% Critical)", "J12", "damping_percent"),
        ("I13", "Dynamic Friction Factor", "J13", "dynamic_fric"),
        ("I14", "Natural Frequency", "J14", "nat_freq"),
    ],
    "jc": [
        ("I19", "Step QTY", "J19", "step_qty"),
        ("I20", "Sample Frequency", "J20", "sample_freq"),
        ("I21", "Physical Length", "J21", "physical_length"),
        ("I22", "EI", "J22", "flexural_rigidity_ei"),
        ("I23", "Celerity", "J23", "celerity"),
        ("I24", "Mass/Physical Length", "J24", "mass_per_unit_len"),
        ("I25", "Fluid Density", "J25", "fluid_densty"),
        ("I26", "Fluid Area", "J26", "fluid_area"),
        ("I27", "Damping (% Critical)", "J27", "damping_percent"),
        ("I28", "Dynamic Friction Factor", "J28", "dynamic_fric"),
        ("I29", "Natural Frequency", "J29", "nat_freq"),
    ],
}
EXAMPLE_INPUT_RANGES = {"pressure": ("B", 3, 7), "lumped": ("E", 3, 27), "jc": ("E", 3, 27)}


def match_case(formula):
    """(solver, match) for a case-table formula, or (None, None)."""
    for solver, rx in (("lumped", LUMPED_RE), ("jc", JC_RE)):
        m = rx.match(formula)
        if m:
            return solver, m
    return None, None


def fmt(v):
    return repr(float(v)) if isinstance(v, float) else str(v)


def cell_range_values(ws, col, r0, r1):
    """Numeric column slice ``{col}{r0}:{col}{r1}`` as a list of 1-element rows (``Range.Value``)."""
    out = []
    for r in range(r0, r1 + 1):
        v = ws[f"{col}{r}"].value
        if not isinstance(v, (int, float)):
            raise SystemExit(f"{ws.title}!{col}{r} is not numeric: {v!r}")
        out.append([v])
    return out


def read_params(ws_val, ws_fml, spec, sheet):
    params = {}
    for label_cell, label, value_cell, name in spec:
        got = ws_fml[label_cell].value
        if got != label:
            raise SystemExit(f"{sheet}!{label_cell} is {got!r}, expected {label!r}")
        v = ws_val[value_cell].value
        if not isinstance(v, (int, float)):
            raise SystemExit(f"{sheet}!{value_cell} is not numeric: {v!r}")
        params[name] = v
    return params


def extract_example(wb_val, wb_fml, out_dir: Path) -> dict:
    """Write ``example/cases.json`` from the ``Example`` sheet; return its manifest entry."""
    ev, ef = wb_val[EXAMPLE_SHEET], wb_fml[EXAMPLE_SHEET]
    cases = []
    for cell, spill, expected_formula, solver in EXAMPLE_FORMULAS:
        value = ef[cell].value
        text = getattr(value, "text", None)
        if text is None or text.replace(" ", "") != expected_formula:
            raise SystemExit(f"{EXAMPLE_SHEET}!{cell}: unexpected formula {text!r}")
        if value.ref != spill:
            raise SystemExit(
                f"{EXAMPLE_SHEET}!{cell}: spill range {value.ref!r}, expected {spill!r}"
            )
        col, r0, r1 = EXAMPLE_INPUT_RANGES[solver]
        inputs = cell_range_values(ev, col, r0, r1)
        m = re.fullmatch(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", spill)
        expected = [
            ev[f"{m.group(1)}{r}"].value for r in range(int(m.group(2)), int(m.group(4)) + 1)
        ]
        if any(not isinstance(v, (int, float)) for v in expected):
            raise SystemExit(f"{EXAMPLE_SHEET}!{spill}: cached outputs missing")
        kwargs = read_params(ev, ef, EXAMPLE_PARAMS[solver], EXAMPLE_SHEET)
        case = {
            "case_id": f"example_{solver}",
            "solver": solver,
            "source": {
                "sheet": EXAMPLE_SHEET,
                "formula_cell": cell,
                "spill": spill,
                "formula": text,
                "input_range": f"{col}{r0}:{col}{r1}",
            },
        }
        if solver == "pressure":
            pipe_id = kwargs.pop("pipe_id_mm")
            # Excel evaluates PI()*(J39^2/4) with PI() = the IEEE double pi (verified against
            # the cached outputs: 10/121 bit-exact vs 3/121 with the 15-digit literal).
            kwargs["pipe_area"] = math.pi * (float(pipe_id) ** 2 / 4)
            case["source"]["pipe_area"] = f"PI()*(J39^2/4) with J39 = {pipe_id} mm"
            case["valve_pattern"] = inputs
        else:
            kwargs["step_qty"] = int(kwargs["step_qty"])
            case["force_history"] = inputs
            case["output_names"] = list(OUTPUT_NAMES)
        case["kwargs"] = kwargs
        case["vba_defaults_used"] = (
            [] if solver == "pressure" else ["restrain_freq_ratio", "gap", "output_type"]
        )
        case["expected"] = expected
        cases.append(case)
    out_dir.mkdir(parents=True, exist_ok=True)
    with open(out_dir / "cases.json", "w") as fh:
        json.dump(cases, fh, indent=1)
        fh.write("\n")
    print(f"wrote {len(cases)} Example cases to {out_dir}", file=sys.stderr)
    return {
        "dir": out_dir.name,
        "files": ["cases.json"],
        "sheet": EXAMPLE_SHEET,
        "cases": len(cases),
        "solvers": {c["solver"]: 1 for c in cases},
        "markers": {},
        "note": "fast cases (milliseconds); run in the default suite",
    }


def extract_curve_10in(wb_val, wb_fml, out_dir: Path, workbook: Path) -> dict:
    """Write the ``curve_10in`` fixture set; return its manifest entry."""
    cv, cf = wb_val[CASE_SHEET], wb_fml[CASE_SHEET]

    # --- common parameters -------------------------------------------------------
    params = {}
    for label_cell, label, value_cell, name in PARAM_CELLS:
        got = cf[label_cell].value
        if got != label:
            raise SystemExit(f"{CASE_SHEET}!{label_cell} is {got!r}, expected {label!r}")
        params[name] = cv[value_cell].value
    if params["step_qty"] != int(params["step_qty"]):
        raise SystemExit("step_qty is not an integer")
    params["step_qty"] = int(params["step_qty"])
    # Excel evaluates $H$14*$H$15 in the formula; reproduce the single double product.
    params["flexural_rigidity_ei"] = float(params.pop("youngs_modulus")) * float(
        params.pop("second_moment")
    )
    force_correction = params.pop("force_correction")

    # --- cases -------------------------------------------------------------------
    cases = []
    force_rows = None
    for r in range(FIRST_CASE_ROW, cf.max_row + 1):
        cell = cf.cell(r, 8)  # H
        text = getattr(cell.value, "text", None)
        if text is None:
            continue  # empty, or a pasted number: not a live case
        solver, m = match_case(text)
        if m is None:
            raise SystemExit(f"{CASE_SHEET}!H{r}: unexpected formula shape: {text}")
        # The five outputs are resultArray(0..4) spilled into H..L; make that explicit.
        if getattr(cell.value, "ref", None) != f"H{r}:L{r}":
            raise SystemExit(
                f"{CASE_SHEET}!H{r}: array formula range is {cell.value.ref!r}, expected 'H{r}:L{r}'"
            )
        c0, c1, e_row, g_row, f_row = (int(x) for x in m.groups())
        if not (e_row == g_row == f_row == r):
            raise SystemExit(f"{CASE_SHEET}!H{r}: formula refers to another row: {text}")
        if force_rows is None:
            force_rows = (c0, c1)
        elif force_rows != (c0, c1):
            raise SystemExit(
                f"{CASE_SHEET}!H{r}: force range differs from row {cases[0]['source_row']}"
            )
        expected = [cv.cell(r, c).value for c in range(8, 13)]  # H..L
        if any(not isinstance(v, (int, float)) for v in expected):
            raise SystemExit(f"{CASE_SHEET}!H{r}: cached outputs missing: {expected}")
        L, f, fric = (cv.cell(r, c).value for c in (5, 6, 7))
        cases.append(
            {
                "case_id": f"r{r:03d}_{solver}_L{L:g}_f{f:g}",
                "source_row": r,
                "solver": solver,
                "physical_length": L,
                "nat_freq": f,
                "dynamic_fric": fric,
                **dict(zip(OUTPUT_NAMES, expected, strict=True)),
            }
        )
    if not cases:
        raise SystemExit("no cases found")

    # --- force history -------------------------------------------------------------
    c0, c1 = force_rows
    n_samples = c1 - c0 + 1
    for r in (c0, c1):
        ref = cf.cell(r, 3).value  # curve_10in!C{r}
        if ref != f"={FORCE_SHEET}!D{r + 1}":
            raise SystemExit(f"{CASE_SHEET}!C{r} is {ref!r}, expected '={FORCE_SHEET}!D{r + 1}'")
    first_force_row = c0 + 1  # Sheet1 row holding sample 1
    ws = wb_val[FORCE_SHEET]
    header = tuple(ws.cell(2, c).value for c in (2, 3, 4))
    if header != FORCE_HEADER:
        raise SystemExit(f"{FORCE_SHEET} header is {header}, expected {FORCE_HEADER}")
    force_rows_out = []
    for r in range(3, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in (2, 3, 4)]
        if all(isinstance(v, (int, float)) for v in vals):
            force_rows_out.append(vals)
        elif any(v is not None for v in vals):
            raise SystemExit(f"{FORCE_SHEET} row {r}: non-numeric data {vals}")
    if first_force_row != 3:
        raise SystemExit(f"force range does not start at {FORCE_SHEET} row 3")
    if n_samples > len(force_rows_out):
        raise SystemExit("force range longer than Sheet1")

    params.update(
        {
            "output_type": 0,
            "n_samples": n_samples,
            "force": {
                "file": "force_history.csv",
                "column": FORCE_COLUMN,
                "source": f"{FORCE_SHEET}!D{first_force_row}:D{first_force_row + n_samples - 1}",
                "excel_expression": f"{CASE_SHEET}!$C${c0}:$C${c1}*1000*$N$13",
                # applied left to right, as Excel evaluates the array expression
                "multipliers": [1000, force_correction],
            },
            "source_workbook": str(workbook.relative_to(REPO)),
            "case_sheet": CASE_SHEET,
            "expected_outputs": list(OUTPUT_NAMES),
        }
    )

    # --- write -----------------------------------------------------------------------
    out_dir.mkdir(parents=True, exist_ok=True)
    with open(out_dir / "force_history.csv", "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(FORCE_HEADER)
        for vals in force_rows_out:
            w.writerow(fmt(v) for v in vals)
    with open(out_dir / "cases.csv", "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(cases[0]))
        w.writeheader()
        for case in cases:
            w.writerow({k: fmt(v) for k, v in case.items()})
    with open(out_dir / "params.json", "w") as fh:
        json.dump(params, fh, indent=2)
        fh.write("\n")

    n_l = sum(c["solver"] == "lumped" for c in cases)
    n_j = sum(c["solver"] == "jc" for c in cases)
    print(
        f"wrote {len(force_rows_out)} force rows (n_samples={n_samples}), "
        f"{len(cases)} cases ({n_l} lumped, {n_j} jc) to {out_dir}",
        file=sys.stderr,
    )
    return {
        "dir": out_dir.name,
        "files": ["params.json", "force_history.csv", "cases.csv"],
        "sheet": CASE_SHEET,
        "cases": len(cases),
        "solvers": {"lumped": n_l, "jc": n_j},
        "markers": {"lumped": "slow", "jc": "jc_slow"},
        "n_samples": n_samples,
    }


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT, help="fixtures root (tests/fixtures)")
    args = ap.parse_args(argv)

    print(f"reading {args.workbook} (values) ...", file=sys.stderr)
    wb_val = openpyxl.load_workbook(args.workbook, data_only=True)
    print("reading formulas ...", file=sys.stderr)
    wb_fml = openpyxl.load_workbook(args.workbook, data_only=False)

    manifest = {
        "generated_by": "tools/extract_fixtures.py",
        "source_workbook": str(args.workbook.relative_to(REPO)),
        "sets": {
            "curve_10in": extract_curve_10in(
                wb_val, wb_fml, args.out / "curve_10in", args.workbook
            ),
            "example": extract_example(wb_val, wb_fml, args.out / "example"),
        },
    }
    with open(args.out / "manifest.json", "w") as fh:
        json.dump(manifest, fh, indent=1)
        fh.write("\n")
    print(f"wrote {args.out / 'manifest.json'}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
